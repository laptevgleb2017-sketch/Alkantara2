import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os
import json
import shutil
from datetime import datetime

try:
    import xlrd
    HAS_XLRD = True
except ImportError:
    HAS_XLRD = False

try:
    import matplotlib
    matplotlib.use('TkAgg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
    HAS_MPL = True
except ImportError:
    HAS_MPL = False


# Русские названия полей для массового редактирования
FIELD_LABELS = {
    'Счёт':           'account',
    'Ответственный':  'responsible',
    'Место хранения': 'location',
}


class AssetManager:
    def __init__(self, root):
        self.root = root
        self.root.title("Ведомость остатков ОС, НМА, НПА")

        desktop = os.path.join(os.path.expanduser('~'), 'Desktop')
        self.app_dir = os.path.join(desktop, 'Учёт')
        self.backup_dir = os.path.join(self.app_dir, 'Резервные копии')
        self.reports_dir = os.path.join(self.app_dir, 'Отчёты')
        for d in [self.app_dir, self.backup_dir, self.reports_dir]:
            os.makedirs(d, exist_ok=True)

        self.settings_path = os.path.join(self.app_dir, 'settings.json')
        self.dict_path = os.path.join(self.app_dir, 'dictionaries.json')
        self.journal_path = os.path.join(self.app_dir, 'journal.log')

        self.assets = []
        self.tree_items = {}
        self.filter_combos = {}
        self.current_file = None
        self.is_report_format = False
        self.sort_field = None
        self.sort_reverse = False
        self.dark_mode = False
        self.window_geometry = '1500x820'
        self.column_widths = {}
        self.autosave_minutes = 5
        self.show_warranty_days = 30

        self.load_settings()
        self.load_dictionaries()

        default_path = os.path.join(self.app_dir, 'assets.xlsx')
        if self.current_file and os.path.exists(self.current_file):
            self.load_file(self.current_file)
        else:
            if not os.path.exists(default_path):
                self.create_empty_file(default_path)
            self.load_file(default_path)

        self.setup_ui()
        self.apply_theme()
        self.refresh()

        self.root.geometry(self.window_geometry)
        self.schedule_autosave()
        self.root.after(1500, self.startup_checks)
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    # ================= НАСТРОЙКИ =================
    def load_settings(self):
        try:
            with open(self.settings_path, 'r', encoding='utf-8') as f:
                s = json.load(f)
            self.current_file = s.get('last_file')
            self.dark_mode = s.get('dark_mode', False)
            self.window_geometry = s.get('window_geometry', '1500x820')
            self.column_widths = s.get('column_widths', {})
            self.autosave_minutes = s.get('autosave_minutes', 5)
            self.show_warranty_days = s.get('show_warranty_days', 30)
        except Exception:
            pass

    def save_settings(self):
        try:
            with open(self.settings_path, 'w', encoding='utf-8') as f:
                json.dump({
                    'last_file': self.current_file,
                    'dark_mode': self.dark_mode,
                    'window_geometry': self.window_geometry,
                    'column_widths': self.get_column_widths(),
                    'autosave_minutes': self.autosave_minutes,
                    'show_warranty_days': self.show_warranty_days,
                }, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def get_column_widths(self):
        try:
            return {c: self.tree.column(c, 'width') for c in self.tree['columns']}
        except Exception:
            return {}

    # ================= СПРАВОЧНИКИ =================
    def load_dictionaries(self):
        defaults = {'accounts': [], 'responsible': [], 'locations': [], 'units': ['шт.']}
        try:
            with open(self.dict_path, 'r', encoding='utf-8') as f:
                self.dictionaries = json.load(f)
        except Exception:
            self.dictionaries = defaults
            self.save_dictionaries()

    def save_dictionaries(self):
        try:
            with open(self.dict_path, 'w', encoding='utf-8') as f:
                json.dump(self.dictionaries, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def update_dict_from_assets(self):
        for a in self.assets:
            for key, field in [('accounts', 'account'),
                               ('responsible', 'responsible'),
                               ('locations', 'location')]:
                val = a.get(field, '').strip()
                if val and val not in self.dictionaries[key]:
                    self.dictionaries[key].append(val)
        self.save_dictionaries()

    # ================= ЖУРНАЛ =================
    def log_action(self, action, details=''):
        try:
            with open(self.journal_path, 'a', encoding='utf-8') as f:
                ts = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                f.write(f"[{ts}] {action}: {details}\n")
        except Exception:
            pass

    # ================= РЕЗЕРВНОЕ КОПИРОВАНИЕ =================
    def create_backup(self):
        if not self.current_file or not os.path.exists(self.current_file):
            return
        ts = datetime.now().strftime('%Y-%m-%d_%H-%M-%S')
        name = f"{os.path.splitext(os.path.basename(self.current_file))[0]}_{ts}.xlsx"
        try:
            shutil.copy2(self.current_file, os.path.join(self.backup_dir, name))
            backups = sorted([f for f in os.listdir(self.backup_dir) if f.endswith('.xlsx')])
            while len(backups) > 30:
                os.remove(os.path.join(self.backup_dir, backups.pop(0)))
        except Exception as e:
            print(f"Backup failed: {e}")

    def restore_backup(self):
        backups = sorted([f for f in os.listdir(self.backup_dir) if f.endswith('.xlsx')],
                         reverse=True)
        if not backups:
            messagebox.showinfo("Резерв", "Резервных копий пока нет")
            return
        dlg = tk.Toplevel(self.root)
        dlg.title("Восстановить из резервной копии")
        dlg.geometry("500x400")
        dlg.grab_set()
        tk.Label(dlg, text="Выберите копию:", font=('Segoe UI', 11)).pack(pady=8)
        lb = tk.Listbox(dlg, width=70, height=15)
        lb.pack(padx=10, pady=5, fill=tk.BOTH, expand=True)
        for b in backups:
            lb.insert('end', b)

        def do_restore():
            sel = lb.curselection()
            if not sel:
                return
            src = os.path.join(self.backup_dir, lb.get(sel[0]))
            if messagebox.askyesno("Подтверждение",
                    f"Восстановить из {os.path.basename(src)}?"):
                if not self.current_file:
                    return
                shutil.copy2(src, self.current_file)
                self.load_file(self.current_file)
                self.refresh()
                self.log_action('RESTORE_BACKUP', src)
                dlg.destroy()
                messagebox.showinfo("Готово", "Восстановлено")

        ttk.Button(dlg, text="Восстановить", command=do_restore).pack(pady=10)

    # ================= ФАЙЛЫ =================
    def create_empty_file(self, path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ведомость"
        headers = ['Счет', 'Ответственный', 'Место хранения', '№ п/п',
                   'Основное средство', 'Инвентарный номер', 'Дата принятия',
                   'Балансовая стоимость', 'Количество', 'Сумма амортизации',
                   'Дата выбытия', 'Причина выбытия', 'Гарантия до',
                   'Следующее ТО']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='2196F3', end_color='2196F3',
                                    fill_type='solid')
        widths = [28, 28, 28, 8, 40, 18, 15, 18, 12, 18, 15, 20, 15, 15]
        for col, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = w
        wb.save(path)

    def detect_format(self, ws):
        for row in ws.iter_rows(min_row=1, max_row=15, max_col=2, values_only=True):
            if not row:
                continue
            if row[0] is None:
                continue
            s = str(row[0])
            if s.strip() == 'Счет':
                return 'flat'
            if 'Ведомость остатков' in s:
                return 'report'
        return 'flat'

    def _ask_sheet(self, names):
        if not names:
            return None
        if len(names) == 1:
            return names[0]
        dlg = tk.Toplevel(self.root)
        dlg.title("Выбор листа")
        dlg.geometry("340x180")
        dlg.grab_set()
        tk.Label(dlg, text="В файле несколько листов.\nВыберите нужный:",
                 font=('Segoe UI', 10)).pack(pady=10)
        var = tk.StringVar(value=names[0])
        ttk.Combobox(dlg, textvariable=var, values=names,
                     state='readonly', width=35).pack(pady=5)
        result = {'ok': False, 'value': None}

        def ok():
            result['ok'] = True
            result['value'] = var.get()
            dlg.destroy()

        ttk.Button(dlg, text="Открыть", command=ok).pack(pady=10)
        dlg.bind('<Return>', lambda e: ok())
        self.root.wait_window(dlg)
        return result['value'] if result['ok'] else None

    def open_file(self):
        path = filedialog.askopenfilename(
            title="Открыть ведомость",
            initialdir=self.app_dir,
            filetypes=[("Excel", "*.xlsx *.xls"),
                       ("Современный Excel", "*.xlsx"),
                       ("Старый Excel", "*.xls"),
                       ("Все файлы", "*.*")])
        if not path:
            return
        self.load_file(path)
        self.refresh()

    def load_file(self, path, sheet_name=None):
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == '.xls':
                self._load_xls(path, sheet_name)
            else:
                wb = openpyxl.load_workbook(path, data_only=True)
                if sheet_name is None and len(wb.sheetnames) > 1:
                    sheet_name = self._ask_sheet(wb.sheetnames)
                    if sheet_name is None:
                        return
                ws = wb[sheet_name] if sheet_name else wb.active
                fmt = self.detect_format(ws)
                if fmt == 'flat':
                    self._load_flat(ws)
                    self.is_report_format = False
                else:
                    self._load_report(ws)
                    self.is_report_format = True
            self.current_file = path
            self.save_settings()
            self.update_title()
            self.log_action('LOAD_FILE', path)
            self.update_dict_from_assets()
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось открыть файл: {e}")

    def _load_xls(self, path, sheet_name=None):
        if not HAS_XLRD:
            messagebox.showerror("Ошибка",
                "Для чтения файлов .xls установите библиотеку xlrd.\n"
                "В командной строке выполните:\npip install xlrd==2.0.1")
            return
        book = xlrd.open_workbook(path)
        names = book.sheet_names()
        if sheet_name is None:
            sheet_name = self._ask_sheet(names)
            if sheet_name is None:
                return
        sheet = book.sheet_by_name(sheet_name)
        self.assets = []
        for r in range(1, sheet.nrows):
            row = [sheet.cell_value(r, c) for c in range(sheet.ncols)]
            if not row or len(row) < 5 or not row[4]:
                continue
            a = self._row_to_asset(row)
            if a:
                self.assets.append(a)
        self.is_report_format = False

    def update_title(self):
        name = os.path.basename(self.current_file) if self.current_file else 'без файла'
        self.root.title(f"Ведомость — {name}")

    def _load_flat(self, ws):
        self.assets = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 5 or not row[4]:
                continue
            a = self._row_to_asset(row)
            if a:
                self.assets.append(a)

    def _row_to_asset(self, row):
        try:
            def s(i):
                return str(row[i]).strip() if i < len(row) and row[i] is not None else ''
            def f(i):
                try:
                    return float(row[i]) if i < len(row) and row[i] not in (None, '') else 0.0
                except (TypeError, ValueError):
                    return 0.0
            def n(i):
                try:
                    return int(float(row[i])) if i < len(row) and row[i] not in (None, '') else 1
                except (TypeError, ValueError):
                    return 1
            return {
                'account': s(0), 'responsible': s(1), 'location': s(2),
                'num': n(3) if n(3) else len(self.assets) + 1,
                'name': s(4), 'inventory': s(5), 'date': s(6),
                'cost': f(7), 'quantity': n(8), 'depreciation': f(9),
                'disposal_date': s(10), 'disposal_reason': s(11),
                'warranty_to': s(12), 'next_to': s(13),
                'residual': f(7) - f(9),
            }
        except Exception:
            return None

    def _load_report(self, ws):
        self.assets = []
        cur_acc = ''; cur_resp = ''; cur_loc = ''; num_counter = 0
        for row in ws.iter_rows(min_row=1, values_only=True):
            if not row: continue
            a = row[0]
            if a is None: continue
            a_str = str(a).strip()
            if not a_str or a_str == 'Итого': continue
            if a_str[:1].isdigit() and ',' in a_str and '.' in a_str.split(',')[0]:
                cur_acc = a_str; continue
            if a_str.isdigit() and len(a_str) >= 15: continue
            c = row[2] if len(row) > 2 else None
            c_empty = (c is None) or (isinstance(c, str) and not c.strip())
            if c_empty:
                if a_str in ('1', '2', '3', '4'): continue
                if any(w in a_str for w in ('МБОУ', 'МКОУ', 'СОШ', 'Гимназия',
                                             'Лицей', 'сад', 'школа', 'д/с')):
                    cur_loc = a_str
                elif any(w in a_str for w in ('вич', 'вна', 'Андреевна',
                                               'Сергеевна', 'Николаевна')):
                    cur_resp = a_str
                continue
            m = row[12] if len(row) > 12 else None
            n = row[13] if len(row) > 13 else None
            o = row[14] if len(row) > 14 else None
            cost = float(m) if isinstance(m, (int, float)) else 0.0
            qty = int(n) if isinstance(n, (int, float)) else 1
            dep = float(o) if isinstance(o, (int, float)) else 0.0
            if isinstance(a, (int, float)):
                num = int(a); num_counter = num
            else:
                num_counter += 1; num = num_counter
            self.assets.append({
                'account': cur_acc, 'responsible': cur_resp, 'location': cur_loc,
                'num': num,
                'name': c.strip() if isinstance(c, str) else str(c),
                'inventory': str(row[8]).strip() if len(row) > 8 and row[8] else '',
                'date': str(row[9]).strip() if len(row) > 9 and row[9] else '',
                'cost': cost, 'quantity': qty, 'depreciation': dep,
                'disposal_date': '', 'disposal_reason': '',
                'warranty_to': '', 'next_to': '',
                'residual': cost - dep,
            })

    def save_file(self):
        if not self.current_file:
            return self.save_file_as()
        if self.current_file.lower().endswith('.xls'):
            messagebox.showwarning("Внимание",
                "Формат .xls не поддерживает запись.\n"
                "Нажмите «Сохранить как» и выберите формат .xlsx.")
            return self.save_file_as()
        try:
            self.create_backup()
            if self.is_report_format:
                self._save_report(self.current_file)
            else:
                self._save_flat(self.current_file)
            self.log_action('SAVE', self.current_file)
            return True
        except PermissionError:
            messagebox.showwarning("Внимание",
                "Файл занят (возможно, открыт в Excel).\nЗакройте Excel и повторите.")
            return False
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить: {e}")
            return False

    def save_file_as(self):
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", initialdir=self.app_dir,
            filetypes=[("Excel 2007+", "*.xlsx"), ("Все файлы", "*.*")])
        if not path:
            return False
        self.current_file = path
        self.is_report_format = False
        self.save_settings()
        self.update_title()
        return self.save_file()

    def _save_flat(self, path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ведомость"
        headers = ['Счет', 'Ответственный', 'Место хранения', '№ п/п',
                   'Основное средство', 'Инвентарный номер', 'Дата принятия',
                   'Балансовая стоимость', 'Количество', 'Сумма амортизации',
                   'Дата выбытия', 'Причина выбытия', 'Гарантия до',
                   'Следующее ТО']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='2196F3', end_color='2196F3',
                                    fill_type='solid')
        for r, a in enumerate(self.assets, 2):
            ws.cell(row=r, column=1, value=a['account'])
            ws.cell(row=r, column=2, value=a['responsible'])
            ws.cell(row=r, column=3, value=a['location'])
            ws.cell(row=r, column=4, value=a['num'])
            ws.cell(row=r, column=5, value=a['name'])
            ws.cell(row=r, column=6, value=a['inventory'])
            ws.cell(row=r, column=7, value=a['date'])
            ws.cell(row=r, column=8, value=a['cost'])
            ws.cell(row=r, column=9, value=a['quantity'])
            ws.cell(row=r, column=10, value=a['depreciation'])
            ws.cell(row=r, column=11, value=a.get('disposal_date', ''))
            ws.cell(row=r, column=12, value=a.get('disposal_reason', ''))
            ws.cell(row=r, column=13, value=a.get('warranty_to', ''))
            ws.cell(row=r, column=14, value=a.get('next_to', ''))
        widths = [28, 28, 28, 8, 40, 18, 15, 18, 12, 18, 15, 20, 15, 15]
        for col, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.freeze_panes = 'A2'
        wb.save(path)

    def _save_report(self, path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Лист_1"
        ws.cell(row=2, column=1,
                value='Ведомость остатков ОС, НМА, НПА').font = Font(bold=True, size=14)
        for c, h in [(13, 'Балансовая стоимость'), (14, 'Количество'),
                     (15, 'Сумма амортизации'), (16, 'Остаточная стоимость')]:
            ws.cell(row=4, column=c, value=h).font = Font(bold=True)
        ws.cell(row=9, column=1, value='№ п/п').font = Font(bold=True)
        ws.cell(row=9, column=3, value='Основное средство').font = Font(bold=True)
        ws.cell(row=9, column=9, value='Инвентарный номер').font = Font(bold=True)
        ws.cell(row=9, column=10, value='Дата принятия к учету').font = Font(bold=True)

        groups = {}
        for a in self.assets:
            groups.setdefault(a['account'], {}) \
                  .setdefault(a['responsible'], {}) \
                  .setdefault(a['location'], []).append(a)

        row = 10
        for acc in sorted(groups.keys()):
            ws.cell(row=row, column=1, value=acc).font = Font(bold=True)
            row += 1
            for resp in sorted(groups[acc].keys()):
                ws.cell(row=row, column=1, value=resp); row += 1
                for loc in sorted(groups[acc][resp].keys()):
                    ws.cell(row=row, column=1, value=loc); row += 1
                    items = groups[acc][resp][loc]
                    for a in sorted(items, key=lambda x: x['num']):
                        ws.cell(row=row, column=1, value=a['num'])
                        ws.cell(row=row, column=3, value=a['name'])
                        ws.cell(row=row, column=9, value=a['inventory'])
                        ws.cell(row=row, column=10, value=a['date'])
                        ws.cell(row=row, column=13, value=a['cost'])
                        ws.cell(row=row, column=14, value=a['quantity'])
                        ws.cell(row=row, column=15, value=a['depreciation'])
                        ws.cell(row=row, column=16, value=a['residual'])
                        row += 1
                    ws.cell(row=row, column=13, value=sum(x['cost'] for x in items))
                    ws.cell(row=row, column=14, value=sum(x['quantity'] for x in items))
                    ws.cell(row=row, column=15, value=sum(x['depreciation'] for x in items))
                    ws.cell(row=row, column=16, value=sum(x['residual'] for x in items))
                    row += 1
        ws.cell(row=row, column=1, value='Итого').font = Font(bold=True)
        ws.cell(row=row, column=13, value=sum(a['cost'] for a in self.assets)).font = Font(bold=True)
        ws.cell(row=row, column=14, value=sum(a['quantity'] for a in self.assets)).font = Font(bold=True)
        ws.cell(row=row, column=15, value=sum(a['depreciation'] for a in self.assets)).font = Font(bold=True)
        ws.cell(row=row, column=16, value=sum(a['residual'] for a in self.assets)).font = Font(bold=True)
        wb.save(path)

    # ================= UI =================
    def setup_ui(self):
        menubar = tk.Menu(self.root)

        m_file = tk.Menu(menubar, tearoff=0)
        m_file.add_command(label="Открыть...", command=self.open_file, accelerator="Ctrl+O")
        m_file.add_command(label="Сохранить", command=self.save_file, accelerator="Ctrl+S")
        m_file.add_command(label="Сохранить как...", command=self.save_file_as, accelerator="Ctrl+Shift+S")
        m_file.add_separator()
        m_file.add_command(label="Восстановить из резервной копии", command=self.restore_backup)
        m_file.add_separator()
        m_file.add_command(label="Выход", command=self.on_close)
        menubar.add_cascade(label="Файл", menu=m_file)

        m_edit = tk.Menu(menubar, tearoff=0)
        m_edit.add_command(label="Добавить", command=self.add_asset, accelerator="Ctrl+N")
        m_edit.add_command(label="Изменить", command=self.edit_asset, accelerator="Ctrl+E")
        m_edit.add_command(label="Удалить", command=self.delete_asset, accelerator="Del")
        m_edit.add_separator()
        m_edit.add_command(label="Массовое изменение", command=self.mass_edit)
        m_edit.add_separator()
        m_edit.add_command(label="Справочники", command=self.show_dictionaries)
        m_edit.add_command(label="Пересчитать амортизацию", command=self.recalc_depreciation)
        menubar.add_cascade(label="Правка", menu=m_edit)

        m_rep = tk.Menu(menubar, tearoff=0)
        m_rep.add_command(label="Ведомость по счёту", command=self.report_by_account)
        m_rep.add_command(label="Отчёт по амортизации", command=self.report_depreciation)
        m_rep.add_command(label="Оборотно-сальдовая", command=self.report_turnover)
        m_rep.add_command(label="Выбыло за период", command=self.report_disposals)
        m_rep.add_separator()
        m_rep.add_command(label="Диаграммы", command=self.show_charts)
        menubar.add_cascade(label="Отчёты", menu=m_rep)

        m_srv = tk.Menu(menubar, tearoff=0)
        m_srv.add_command(label="Проверить целостность", command=self.check_integrity)
        m_srv.add_command(label="Показать напоминания", command=self.show_notifications)
        m_srv.add_command(label="Открыть журнал", command=self.open_journal)
        menubar.add_cascade(label="Сервис", menu=m_srv)

        m_help = tk.Menu(menubar, tearoff=0)
        m_help.add_command(label="Горячие клавиши", command=self.show_help)
        m_help.add_command(label="О программе", command=lambda: messagebox.showinfo(
            "О программе", "Учёт имущества\n© Ольгерд"))
        menubar.add_cascade(label="Справка", menu=m_help)
        self.root.config(menu=menubar)

        self.top = tk.Frame(self.root, bg='#f0f0f0')
        self.top.pack(fill=tk.X, padx=15, pady=(10, 5))
        tk.Label(self.top, text="📊 Ведомость ОС, НМА, НПА",
                 font=('Segoe UI', 16, 'bold'),
                 bg='#f0f0f0').pack(side=tk.LEFT)

        btns = tk.Frame(self.top, bg='#f0f0f0')
        btns.pack(side=tk.RIGHT)
        for text, cmd, color in [
            ("📂 Открыть", self.open_file, '#607D8B'),
            ("💾 Сохранить", self.save_file, '#009688'),
            ("➕", self.add_asset, '#4CAF50'),
            ("✏️", self.edit_asset, '#FF9800'),
            ("🗑", self.delete_asset, '#F44336'),
            ("Массово", self.mass_edit, '#795548'),
            ("🌓 Тема", self.toggle_theme, '#455A64'),
        ]:
            tk.Button(btns, text=text, command=cmd, bg=color, fg='white',
                      relief=tk.FLAT, padx=10, pady=6,
                      cursor='hand2').pack(side=tk.LEFT, padx=2)

        self.flt = tk.Frame(self.root, bg='#ffffff', relief=tk.RAISED, bd=1)
        self.flt.pack(fill=tk.X, padx=15, pady=5)
        inner = tk.Frame(self.flt, bg='#ffffff')
        inner.pack(fill=tk.X, padx=10, pady=8)

        self.f_account = tk.StringVar()
        self.f_resp = tk.StringVar()
        self.f_loc = tk.StringVar()
        self.f_search = tk.StringVar()
        self.f_cost_min = tk.StringVar()
        self.f_cost_max = tk.StringVar()
        self.f_date_from = tk.StringVar()
        self.f_date_to = tk.StringVar()
        self.f_only_active = tk.BooleanVar(value=False)
        self.f_full_depreciated = tk.BooleanVar(value=False)

        tk.Label(inner, text="🔍", bg='#ffffff').pack(side=tk.LEFT)
        e = tk.Entry(inner, textvariable=self.f_search, width=22)
        e.pack(side=tk.LEFT, padx=4)
        e.bind('<KeyRelease>', lambda ev: self.refresh())
        self.search_entry = e

        for label, var, key in [("Счёт:", self.f_account, 'account'),
                                 ("Ответств.:", self.f_resp, 'responsible'),
                                 ("Место:", self.f_loc, 'location')]:
            tk.Label(inner, text=label, bg='#ffffff').pack(side=tk.LEFT, padx=(8, 2))
            cb = ttk.Combobox(inner, textvariable=var, width=18, state='readonly')
            cb.pack(side=tk.LEFT)
            cb.bind('<<ComboboxSelected>>', lambda ev: self.refresh())
            self.filter_combos[key] = cb

        ttk.Button(inner, text="▼ Доп.", width=8,
                   command=self.toggle_advanced).pack(side=tk.LEFT, padx=6)
        ttk.Button(inner, text="✖", width=3,
                   command=self.reset_filters).pack(side=tk.LEFT)

        self.adv = tk.Frame(self.root, bg='#f9f9f9', relief=tk.RIDGE, bd=1)
        self.adv_visible = False
        self.adv_inner = tk.Frame(self.adv, bg='#f9f9f9')
        self.adv_inner.pack(fill=tk.X, padx=10, pady=6)
        for label, var, w in [("Стоимость от:", self.f_cost_min, 10),
                               ("до:", self.f_cost_max, 10),
                               ("Дата от:", self.f_date_from, 10),
                               ("до:", self.f_date_to, 10)]:
            tk.Label(self.adv_inner, text=label, bg='#f9f9f9').pack(side=tk.LEFT, padx=2)
            tk.Entry(self.adv_inner, textvariable=var, width=w).pack(side=tk.LEFT)
        tk.Checkbutton(self.adv_inner, text="Только активные",
                       variable=self.f_only_active, bg='#f9f9f9',
                       command=self.refresh).pack(side=tk.LEFT, padx=10)
        tk.Checkbutton(self.adv_inner, text="Полностью самортизированные",
                       variable=self.f_full_depreciated, bg='#f9f9f9',
                       command=self.refresh).pack(side=tk.LEFT, padx=10)
        ttk.Button(self.adv_inner, text="Применить",
                   command=self.refresh).pack(side=tk.LEFT, padx=6)

        table = tk.Frame(self.root, bg='#ffffff', relief=tk.SOLID, bd=1)
        table.pack(fill=tk.BOTH, expand=True, padx=15, pady=5)

        cols = ('num', 'name', 'inventory', 'date',
                'cost', 'qty', 'depreciation', 'residual',
                'warranty', 'next_to', 'disposal')
        heads = ('№ п/п', 'Основное средство', 'Инв. №', 'Дата прин.',
                 'Бал. ст-ть', 'Кол.', 'Аморт.', 'Остат.',
                 'Гарантия до', 'След. ТО', 'Выбытие')
        widths = (55, 320, 130, 95, 120, 55, 120, 120, 100, 100, 100)

        self.tree = ttk.Treeview(table, columns=cols, show='tree headings',
                                  selectmode='extended')
        self.tree.heading('#0', text='Счёт / Ответственный / Место')
        self.tree.column('#0', width=340, anchor='w')
        for c, h, w in zip(cols, heads, widths):
            self.tree.heading(c, text=h, command=lambda cc=c: self.sort_by(cc))
            self.tree.column(c, width=w,
                anchor='center' if c != 'name' else 'w')

        vs = ttk.Scrollbar(table, orient='vertical', command=self.tree.yview)
        hs = ttk.Scrollbar(table, orient='horizontal', command=self.tree.xview)
        self.tree.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        self.tree.grid(row=0, column=0, sticky='nsew')
        vs.grid(row=0, column=1, sticky='ns')
        hs.grid(row=1, column=0, sticky='ew')
        table.grid_rowconfigure(0, weight=1)
        table.grid_columnconfigure(0, weight=1)

        self.tree.bind('<Double-Button-1>', self.on_double_click)
        self.tree.bind('<Button-3>', self.show_context_menu)

        bottom = tk.Frame(self.root, bg='#f0f0f0')
        bottom.pack(fill=tk.X, padx=15, pady=(0, 8))
        self.file_label = tk.Label(bottom, text='', font=('Segoe UI', 9),
                                    bg='#f0f0f0', fg='#555', anchor='w')
        self.file_label.pack(side=tk.LEFT)
        self.autosave_label = tk.Label(bottom, text='', font=('Segoe UI', 9),
                                        bg='#f0f0f0', fg='#2E7D32')
        self.autosave_label.pack(side=tk.LEFT, padx=10)
        self.totals = tk.Label(bottom, text='', font=('Segoe UI', 10, 'bold'),
                                bg='#f0f0f0', anchor='e', fg='#1a237e')
        self.totals.pack(side=tk.RIGHT)

        self.root.bind('<Control-o>', lambda e: self.open_file())
        self.root.bind('<Control-s>', lambda e: self.save_file())
        self.root.bind('<Control-S>', lambda e: self.save_file_as())
        self.root.bind('<Control-f>', lambda e: self.search_entry.focus_set())
        self.root.bind('<Control-n>', lambda e: self.add_asset())
        self.root.bind('<Control-e>', lambda e: self.edit_asset())
        self.root.bind('<Delete>', lambda e: self.delete_asset())
        self.root.bind('<F5>', lambda e: self.refresh())

    def apply_theme(self):
        # Размер шрифта для групп: базовый 9 + 4 = 13
        GRP_FONT_SIZE = 13
        if self.dark_mode:
            bg = '#2b2b2b'; panel = '#3c3c3c'
            self.tree.tag_configure('group1', background='#1e3a5f',
                                     foreground='#ffffff',
                                     font=('Segoe UI', GRP_FONT_SIZE, 'bold'))
            self.tree.tag_configure('group2', background='#2a4a2a',
                                     foreground='#ffffff',
                                     font=('Segoe UI', GRP_FONT_SIZE, 'italic'))
            self.tree.tag_configure('group3', background='#5f4a1e',
                                     foreground='#ffffff',
                                     font=('Segoe UI', GRP_FONT_SIZE))
        else:
            bg = '#f0f0f0'; panel = '#ffffff'
            self.tree.tag_configure('group1', background='#E3F2FD',
                                     font=('Segoe UI', GRP_FONT_SIZE, 'bold'))
            self.tree.tag_configure('group2', background='#F1F8E9',
                                     font=('Segoe UI', GRP_FONT_SIZE, 'italic'))
            self.tree.tag_configure('group3', background='#FFF8E1',
                                     font=('Segoe UI', GRP_FONT_SIZE))

        for w in [self.root, self.top, self.flt]:
            try: w.configure(bg=bg)
            except Exception: pass
        style = ttk.Style()
        style.theme_use('default')
        style.configure('Treeview', rowheight=30, font=('Segoe UI', 9), background=panel)
        style.configure('Treeview.Heading', font=('Segoe UI', 9, 'bold'))

    def toggle_theme(self):
        self.dark_mode = not self.dark_mode
        self.apply_theme()
        self.save_settings()

    def toggle_advanced(self):
        if self.adv_visible:
            self.adv.pack_forget()
        else:
            self.adv.pack(fill=tk.X, padx=15, pady=(0, 5), before=self.tree.master)
        self.adv_visible = not self.adv_visible

    # ================= ФИЛЬТРЫ / СОРТИРОВКА =================
    def reset_filters(self):
        for v in [self.f_account, self.f_resp, self.f_loc, self.f_search,
                  self.f_cost_min, self.f_cost_max, self.f_date_from, self.f_date_to]:
            v.set('')
        self.f_only_active.set(False)
        self.f_full_depreciated.set(False)
        self.refresh()

    def filtered(self):
        acc = self.f_account.get().strip()
        resp = self.f_resp.get().strip()
        loc = self.f_loc.get().strip()
        search = self.f_search.get().strip().lower()
        try: cmin = float(self.f_cost_min.get()) if self.f_cost_min.get() else None
        except ValueError: cmin = None
        try: cmax = float(self.f_cost_max.get()) if self.f_cost_max.get() else None
        except ValueError: cmax = None
        dfrom = self.f_date_from.get().strip()
        dto = self.f_date_to.get().strip()
        only_active = self.f_only_active.get()
        full_dep = self.f_full_depreciated.get()

        out = []
        for a in self.assets:
            if acc and a['account'] != acc: continue
            if resp and a['responsible'] != resp: continue
            if loc and a['location'] != loc: continue
            if cmin is not None and a['cost'] < cmin: continue
            if cmax is not None and a['cost'] > cmax: continue
            if dfrom and a['date']:
                try:
                    if self._parse_date(a['date']) < self._parse_date(dfrom): continue
                except Exception: pass
            if dto and a['date']:
                try:
                    if self._parse_date(a['date']) > self._parse_date(dto): continue
                except Exception: pass
            if only_active and a.get('disposal_date'): continue
            if full_dep and not (a['cost'] > 0 and a['depreciation'] >= a['cost']):
                continue
            if search:
                hay = ' '.join([a['name'], a['inventory'], a['responsible'],
                                a['location'], a['account']]).lower()
                if search not in hay: continue
            out.append(a)

        if self.sort_field:
            out.sort(key=lambda x: self._sort_key(x, self.sort_field),
                     reverse=self.sort_reverse)
        return out

    def _parse_date(self, s):
        for fmt in ('%d.%m.%Y', '%Y-%m-%d', '%d/%m/%Y'):
            try: return datetime.strptime(s.strip(), fmt)
            except Exception: continue
        return datetime(1900, 1, 1)

    def _sort_key(self, a, f):
        if f in ('cost', 'depreciation', 'residual', 'quantity', 'num'):
            return a.get(f, 0)
        if f == 'date':
            return self._parse_date(a.get('date', ''))
        return str(a.get(f, '')).lower()

    def sort_by(self, field):
        if self.sort_field == field:
            self.sort_reverse = not self.sort_reverse
        else:
            self.sort_field = field
            self.sort_reverse = False
        self.refresh()

    def update_filter_values(self):
        for key, field in [('account', 'account'),
                           ('responsible', 'responsible'),
                           ('location', 'location')]:
            values = sorted({a[field] for a in self.assets if a[field]})
            cb = self.filter_combos.get(key)
            if cb is not None:
                current = cb.get()
                cb['values'] = [''] + values
                if current not in values:
                    cb.set('')

    def refresh(self):
        self.update_filter_values()
        for i in self.tree.get_children():
            self.tree.delete(i)
        self.tree_items.clear()

        data = self.filtered()
        groups = {}
        for a in data:
            groups.setdefault(a['account'], {}) \
                  .setdefault(a['responsible'], {}) \
                  .setdefault(a['location'], []).append(a)

        for acc in sorted(groups.keys()):
            acc_items = [x for r in groups[acc].values()
                         for loc in r.values() for x in loc]
            s = self._sum(acc_items)
            acc_id = self.tree.insert('', 'end',
                text=f"📁 {acc or '— без счёта'}",
                values=('', '', '', '', f"{s[0]:,.2f}", s[1],
                        f"{s[2]:,.2f}", f"{s[3]:,.2f}", '', '', ''),
                tags=('group1',), open=True)
            for resp in sorted(groups[acc].keys()):
                resp_items = [x for loc in groups[acc][resp].values() for x in loc]
                s = self._sum(resp_items)
                resp_id = self.tree.insert(acc_id, 'end',
                    text=f"👤 {resp or '— без ответ.'}",
                    values=('', '', '', '', f"{s[0]:,.2f}", s[1],
                            f"{s[2]:,.2f}", f"{s[3]:,.2f}", '', '', ''),
                    tags=('group2',), open=False)
                for loc in sorted(groups[acc][resp].keys()):
                    items = groups[acc][resp][loc]
                    s = self._sum(items)
                    loc_id = self.tree.insert(resp_id, 'end',
                        text=f"📍 {loc or '— без места'}",
                        values=('', '', '', '', f"{s[0]:,.2f}", s[1],
                                f"{s[2]:,.2f}", f"{s[3]:,.2f}", '', '', ''),
                        tags=('group3',), open=False)
                    for a in sorted(items, key=lambda x: x['num']):
                        iid = self.tree.insert(loc_id, 'end', text='', values=(
                            a['num'], a['name'], a['inventory'], a['date'],
                            f"{a['cost']:,.2f}", a['quantity'],
                            f"{a['depreciation']:,.2f}", f"{a['residual']:,.2f}",
                            a.get('warranty_to', ''), a.get('next_to', ''),
                            a.get('disposal_date', '')))
                        self.tree_items[iid] = a

        self.totals.config(text=self._totals_text(data))
        self.file_label.config(text=self.current_file or '')

    @staticmethod
    def _sum(items):
        return (sum(x['cost'] for x in items),
                sum(x['quantity'] for x in items),
                sum(x['depreciation'] for x in items),
                sum(x['residual'] for x in items))

    def _totals_text(self, data):
        c, q, d, r = self._sum(data)
        return (f"Итого: Бал. {c:,.2f} ₽ | Кол. {q} | "
                f"Аморт. {d:,.2f} ₽ | Остат. {r:,.2f} ₽")

    # ================= CRUD =================
    def on_double_click(self, event):
        iid = self.tree.identify_row(event.y)
        if iid in self.tree_items:
            self.show_asset_card(self.tree_items[iid])

    def show_context_menu(self, event):
        iid = self.tree.identify_row(event.y)
        if iid not in self.tree_items:
            return
        self.tree.selection_set(iid)
        asset = self.tree_items[iid]
        m = tk.Menu(self.root, tearoff=0)
        m.add_command(label="Карточка", command=lambda: self.show_asset_card(asset))
        m.add_command(label="Редактировать", command=lambda: self.edit_asset(asset))
        m.add_command(label="Дублировать", command=lambda: self.duplicate_asset(asset))
        m.add_separator()
        m.add_command(label="Списать", command=lambda: self.dispose_asset(asset))
        m.add_separator()
        m.add_command(label="Удалить", command=lambda: self.delete_asset(asset))
        m.tk_popup(event.x_root, event.y_root)

    def add_asset(self):
        dlg = AssetDialog(self.root, "Добавить", dictionaries=self.dictionaries,
                          last=self._last_values())
        self.root.wait_window(dlg.dialog)
        if dlg.result:
            if not self._validate_unique(dlg.result):
                return
            dlg.result['num'] = max([a['num'] for a in self.assets] + [0]) + 1
            dlg.result['residual'] = dlg.result['cost'] - dlg.result['depreciation']
            self.assets.append(dlg.result)
            self.save_file()
            self.log_action('ADD', f"{dlg.result['name']} ({dlg.result['inventory']})")
            self.refresh()

    def _last_values(self):
        if not self.assets:
            return {}
        last = self.assets[-1]
        return {'account': last['account'],
                'responsible': last['responsible'],
                'location': last['location']}

    def _validate_unique(self, asset):
        inv = asset.get('inventory', '').strip()
        if not inv:
            return True
        for a in self.assets:
            if a['inventory'].strip() == inv:
                messagebox.showwarning("Дубликат",
                    f"Инвентарный номер «{inv}» уже используется:\n"
                    f"«{a['name']}» ({a['location']})")
                return False
        return True

    def duplicate_asset(self, asset):
        new = dict(asset)
        new['name'] = asset['name'] + ' (копия)'
        new['inventory'] = ''
        dlg = AssetDialog(self.root, "Дублировать", new,
                          dictionaries=self.dictionaries)
        self.root.wait_window(dlg.dialog)
        if dlg.result:
            if not self._validate_unique(dlg.result):
                return
            dlg.result['num'] = max([a['num'] for a in self.assets] + [0]) + 1
            dlg.result['residual'] = dlg.result['cost'] - dlg.result['depreciation']
            self.assets.append(dlg.result)
            self.save_file()
            self.log_action('DUP', new['name'])
            self.refresh()

    def edit_asset(self, asset=None):
        if asset is None:
            iid = self.tree.focus()
            asset = self.tree_items.get(iid)
        if not asset:
            messagebox.showwarning("Внимание", "Выберите актив")
            return
        dlg = AssetDialog(self.root, "Редактировать", asset,
                          dictionaries=self.dictionaries)
        self.root.wait_window(dlg.dialog)
        if dlg.result:
            if dlg.result['inventory'].strip() != asset['inventory'].strip() \
               and not self._validate_unique(dlg.result):
                return
            dlg.result['num'] = asset['num']
            dlg.result['residual'] = dlg.result['cost'] - dlg.result['depreciation']
            idx = self.assets.index(asset)
            self.assets[idx] = dlg.result
            self.save_file()
            self.log_action('EDIT', f"{asset['name']} ({asset['inventory']})")
            self.refresh()

    def delete_asset(self, asset=None):
        if asset is None:
            iid = self.tree.focus()
            asset = self.tree_items.get(iid)
        if not asset:
            messagebox.showwarning("Внимание", "Выберите актив")
            return
        if messagebox.askyesno("Подтверждение", f"Удалить «{asset['name']}»?"):
            self.assets.remove(asset)
            self.save_file()
            self.log_action('DELETE', f"{asset['name']} ({asset['inventory']})")
            self.refresh()

    def dispose_asset(self, asset):
        dlg = tk.Toplevel(self.root)
        dlg.title("Списание актива")
        dlg.geometry("400x220")
        dlg.grab_set()
        tk.Label(dlg, text=f"Списать: {asset['name']}",
                 font=('Segoe UI', 11, 'bold')).pack(pady=10)
        tk.Label(dlg, text="Дата списания (ДД.ММ.ГГГГ):").pack()
        d_var = tk.StringVar(value=datetime.now().strftime('%d.%m.%Y'))
        tk.Entry(dlg, textvariable=d_var, width=20).pack(pady=4)
        tk.Label(dlg, text="Причина:").pack()
        r_var = tk.StringVar()
        ttk.Combobox(dlg, textvariable=r_var, width=30,
                     values=['Списание по износу', 'Продажа', 'Передача',
                             'Утеря', 'Порча', 'Другое']).pack(pady=4)

        def do():
            asset['disposal_date'] = d_var.get()
            asset['disposal_reason'] = r_var.get()
            self.save_file()
            self.log_action('DISPOSE', f"{asset['name']}: {r_var.get()} {d_var.get()}")
            self.refresh()
            dlg.destroy()

        ttk.Button(dlg, text="Списать", command=do).pack(pady=10)

    def mass_edit(self):
        selected = [self.tree_items[i] for i in self.tree.selection()
                    if i in self.tree_items]
        if not selected:
            messagebox.showwarning("Внимание", "Выберите строки актива")
            return
        dlg = tk.Toplevel(self.root)
        dlg.title("Массовое изменение")
        dlg.geometry("440x260")
        dlg.grab_set()
        tk.Label(dlg, text=f"Выбрано активов: {len(selected)}",
                 font=('Segoe UI', 11, 'bold')).pack(pady=8)
        tk.Label(dlg, text="Что изменить:", bg='#f0f0f0',
                 font=('Segoe UI', 10)).pack(anchor='w', padx=20)
        field_var = tk.StringVar(value='Ответственный')
        ttk.Combobox(dlg, textvariable=field_var, state='readonly', width=30,
                     values=list(FIELD_LABELS.keys())).pack(padx=20, pady=4)
        tk.Label(dlg, text="Новое значение:", bg='#f0f0f0',
                 font=('Segoe UI', 10)).pack(anchor='w', padx=20, pady=(8, 0))
        value_var = tk.StringVar()
        tk.Entry(dlg, textvariable=value_var, width=40,
                 font=('Segoe UI', 10)).pack(padx=20, pady=4)

        def do():
            if not value_var.get().strip():
                messagebox.showwarning("Внимание", "Введите новое значение")
                return
            ru_field = field_var.get()
            en_field = FIELD_LABELS.get(ru_field)
            if not en_field:
                return
            new_val = value_var.get().strip()
            for a in selected:
                a[en_field] = new_val
            self.save_file()
            self.log_action('MASS_EDIT',
                            f"{len(selected)} × {ru_field} = {new_val}")
            self.refresh()
            dlg.destroy()
            messagebox.showinfo("Готово",
                f"Изменено активов: {len(selected)}\n{ru_field} → {new_val}")

        ttk.Button(dlg, text="Применить", command=do).pack(pady=12)

    # ================= КАРТОЧКА =================
    def show_asset_card(self, asset):
        dlg = tk.Toplevel(self.root)
        dlg.title(f"Карточка: {asset['name']}")
        dlg.geometry("560x500")
        dlg.grab_set()
        txt = tk.Text(dlg, wrap=tk.WORD, font=('Segoe UI', 10))
        txt.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        fields = [('Счёт', 'account'), ('Ответственный', 'responsible'),
                  ('Место хранения', 'location'), ('№ п/п', 'num'),
                  ('Наименование', 'name'), ('Инвентарный номер', 'inventory'),
                  ('Дата принятия', 'date'),
                  ('Балансовая стоимость', 'cost'),
                  ('Количество', 'quantity'),
                  ('Сумма амортизации', 'depreciation'),
                  ('Остаточная стоимость', 'residual'),
                  ('Гарантия до', 'warranty_to'),
                  ('Следующее ТО', 'next_to'),
                  ('Дата выбытия', 'disposal_date'),
                  ('Причина выбытия', 'disposal_reason')]
        for label, key in fields:
            val = asset.get(key, '')
            if key in ('cost', 'depreciation', 'residual') and isinstance(val, (int, float)):
                val = f"{val:,.2f} ₽"
            txt.insert(tk.END, f"{label}:\n", 'bold')
            txt.insert(tk.END, f"    {val}\n\n")
        txt.tag_configure('bold', font=('Segoe UI', 10, 'bold'))
        txt.config(state=tk.DISABLED)

    # ================= СПРАВОЧНИКИ / АМОРТИЗАЦИЯ =================
    def show_dictionaries(self):
        dlg = tk.Toplevel(self.root)
        dlg.title("Справочники")
        dlg.geometry("500x500")
        dlg.grab_set()
        nb = ttk.Notebook(dlg)
        nb.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        for name, key in [('Счета', 'accounts'),
                          ('Ответственные', 'responsible'),
                          ('Места хранения', 'locations'),
                          ('Единицы', 'units')]:
            frame = tk.Frame(nb)
            nb.add(frame, text=name)
            lb = tk.Listbox(frame)
            lb.pack(fill=tk.BOTH, expand=True, side=tk.LEFT, padx=5, pady=5)
            for v in self.dictionaries.get(key, []):
                lb.insert('end', v)
            btns = tk.Frame(frame)
            btns.pack(side=tk.RIGHT, fill=tk.Y, padx=5, pady=5)

            def add_item(k=key, l=lb):
                v = simpledialog.askstring("Добавить", "Значение:")
                if v and v not in self.dictionaries[k]:
                    self.dictionaries[k].append(v)
                    l.insert('end', v)
                    self.save_dictionaries()

            def del_item(k=key, l=lb):
                sel = l.curselection()
                if not sel:
                    return
                v = l.get(sel[0])
                l.delete(sel[0])
                self.dictionaries[k].remove(v)
                self.save_dictionaries()

            ttk.Button(btns, text="➕", command=add_item).pack(pady=2, fill=tk.X)
            ttk.Button(btns, text="✖", command=del_item).pack(pady=2, fill=tk.X)

    def recalc_depreciation(self):
        dlg = tk.Toplevel(self.root)
        dlg.title("Пересчёт амортизации")
        dlg.geometry("380x220")
        dlg.grab_set()
        tk.Label(dlg, text="Дата пересчёта (ДД.ММ.ГГГГ):").pack(pady=8)
        d_var = tk.StringVar(value=datetime.now().strftime('%d.%m.%Y'))
        tk.Entry(dlg, textvariable=d_var, width=20).pack()
        tk.Label(dlg, text="Срок службы (мес.):").pack(pady=8)
        m_var = tk.StringVar(value='60')
        tk.Entry(dlg, textvariable=m_var, width=10).pack()

        def do():
            try:
                months = int(m_var.get())
            except ValueError:
                months = 60
            changed = 0
            for a in self.assets:
                if not a.get('date'):
                    continue
                try:
                    start = self._parse_date(a['date'])
                    now = self._parse_date(d_var.get())
                    used = (now.year - start.year) * 12 + (now.month - start.month)
                    if used < 0:
                        used = 0
                    monthly = a['cost'] / months if months else 0
                    new_dep = min(monthly * used, a['cost'])
                    if abs(new_dep - a['depreciation']) > 0.01:
                        a['depreciation'] = round(new_dep, 2)
                        a['residual'] = a['cost'] - a['depreciation']
                        changed += 1
                except Exception:
                    pass
            self.save_file()
            self.log_action('RECALC_DEPR', f"Изменено {changed} записей")
            self.refresh()
            dlg.destroy()
            messagebox.showinfo("Готово", f"Обновлено записей: {changed}")

        ttk.Button(dlg, text="Пересчитать", command=do).pack(pady=12)

    # ================= ОТЧЁТЫ =================
    def report_by_account(self):
        self._printable_report("Ведомость по счёту", self._build_account_report)

    def _build_account_report(self):
        lines = ["ВЕДОМОСТЬ ОСТАТКОВ ОС, НМА, НПА",
                 f"Сформирована: {datetime.now().strftime('%d.%m.%Y %H:%M')}", ""]
        groups = {}
        for a in self.filtered():
            groups.setdefault(a['account'], []).append(a)
        for acc in sorted(groups.keys()):
            items = groups[acc]
            lines.append(f"Счёт: {acc}")
            lines.append("-" * 100)
            lines.append(f"{'№':>4} | {'Наименование':40} | {'Инв.№':15} | "
                         f"{'Бал.':>12} | {'Аморт.':>12} | {'Остат.':>12}")
            lines.append("-" * 100)
            for a in sorted(items, key=lambda x: x['num']):
                lines.append(f"{a['num']:>4} | {a['name'][:40]:40} | "
                             f"{a['inventory'][:15]:15} | "
                             f"{a['cost']:>12,.2f} | {a['depreciation']:>12,.2f} | "
                             f"{a['residual']:>12,.2f}")
            s = self._sum(items)
            lines.append("-" * 100)
            lines.append(f"{'ИТОГО:':>4} | {'':40} | {'':15} | "
                         f"{s[0]:>12,.2f} | {s[2]:>12,.2f} | {s[3]:>12,.2f}")
            lines.append("")
        t = self._sum(self.filtered())
        lines.append("=" * 100)
        lines.append(f"ВСЕГО: {t[0]:,.2f} ₽ | Амортизация: {t[2]:,.2f} ₽ | "
                     f"Остаточная: {t[3]:,.2f} ₽")
        return "\n".join(lines)

    def report_depreciation(self):
        dlg = tk.Toplevel(self.root)
        dlg.title("Амортизация за период")
        dlg.geometry("900x600")
        dlg.grab_set()
        tk.Label(dlg, text="Отчёт по амортизации",
                 font=('Segoe UI', 12, 'bold')).pack(pady=6)
        txt = tk.Text(dlg, font=('Consolas', 9), wrap=tk.NONE)
        txt.pack(fill=tk.BOTH, expand=True, padx=10, pady=6)
        data = self.filtered()
        total = self._sum(data)
        txt.insert(tk.END, f"Дата: {datetime.now().strftime('%d.%m.%Y')}\n\n")
        txt.insert(tk.END, f"{'Счёт':30} | {'Ст-ть':>15} | {'Аморт.':>15} | {'Остат.':>15}\n")
        txt.insert(tk.END, "-" * 85 + "\n")
        groups = {}
        for a in data:
            groups.setdefault(a['account'], []).append(a)
        for acc in sorted(groups.keys()):
            s = self._sum(groups[acc])
            txt.insert(tk.END, f"{acc[:30]:30} | {s[0]:>15,.2f} | "
                               f"{s[2]:>15,.2f} | {s[3]:>15,.2f}\n")
        txt.insert(tk.END, "-" * 85 + "\n")
        txt.insert(tk.END, f"{'ИТОГО':30} | {total[0]:>15,.2f} | "
                           f"{total[2]:>15,.2f} | {total[3]:>15,.2f}\n")
        txt.config(state=tk.DISABLED)

    def report_turnover(self):
        dlg = tk.Toplevel(self.root)
        dlg.title("Оборотно-сальдовая")
        dlg.geometry("800x500")
        dlg.grab_set()
        txt = tk.Text(dlg, font=('Consolas', 10))
        txt.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        txt.insert(tk.END, "ОБОРОТНО-САЛЬДОВАЯ ВЕДОМОСТЬ\n")
        txt.insert(tk.END, f"Дата: {datetime.now().strftime('%d.%m.%Y')}\n\n")
        txt.insert(tk.END, f"{'Счёт':30} | {'Нач.':>12} | {'Поступ.':>12} | "
                           f"{'Выбыло':>12} | {'Кон.':>12}\n")
        txt.insert(tk.END, "-" * 90 + "\n")
        groups = {}
        for a in self.assets:
            groups.setdefault(a['account'], []).append(a)
        for acc in sorted(groups.keys()):
            items = groups[acc]
            active = [x for x in items if not x.get('disposal_date')]
            disposed = [x for x in items if x.get('disposal_date')]
            beginning = sum(x['cost'] for x in items)
            incoming = sum(x['cost'] for x in active)
            outgoing = sum(x['cost'] for x in disposed)
            ending = beginning + incoming - outgoing
            txt.insert(tk.END, f"{acc[:30]:30} | {beginning:>12,.0f} | "
                               f"{incoming:>12,.0f} | {outgoing:>12,.0f} | "
                               f"{ending:>12,.0f}\n")
        txt.config(state=tk.DISABLED)

    def report_disposals(self):
        dlg = tk.Toplevel(self.root)
        dlg.title("Выбыло за период")
        dlg.geometry("900x500")
        dlg.grab_set()
        top = tk.Frame(dlg); top.pack(fill=tk.X, padx=10, pady=6)
        tk.Label(top, text="От:").pack(side=tk.LEFT)
        d1 = tk.StringVar(value='01.01.2020')
        tk.Entry(top, textvariable=d1, width=12).pack(side=tk.LEFT, padx=4)
        tk.Label(top, text="до:").pack(side=tk.LEFT)
        d2 = tk.StringVar(value=datetime.now().strftime('%d.%m.%Y'))
        tk.Entry(top, textvariable=d2, width=12).pack(side=tk.LEFT, padx=4)
        txt = tk.Text(dlg, font=('Consolas', 9))
        txt.pack(fill=tk.BOTH, expand=True, padx=10, pady=6)

        def do():
            txt.delete('1.0', tk.END)
            try:
                f1 = self._parse_date(d1.get()); f2 = self._parse_date(d2.get())
            except Exception:
                return
            items = [a for a in self.assets if a.get('disposal_date') and
                     f1 <= self._parse_date(a['disposal_date']) <= f2]
            txt.insert(tk.END, f"Выбывших активов: {len(items)}\n\n")
            for a in items:
                txt.insert(tk.END, f"{a['disposal_date']} | {a['name'][:40]:40} | "
                                   f"{a['inventory']} | {a['disposal_reason']}\n")

        ttk.Button(top, text="Показать", command=do).pack(side=tk.LEFT, padx=8)

    def _printable_report(self, title, text_builder):
        dlg = tk.Toplevel(self.root)
        dlg.title(title)
        dlg.geometry("1000x650")
        dlg.grab_set()
        txt = tk.Text(dlg, font=('Consolas', 9), wrap=tk.NONE)
        txt.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        txt.insert(tk.END, text_builder())
        txt.config(state=tk.DISABLED)

        def save_as_txt():
            p = filedialog.asksaveasfilename(
                defaultextension=".txt", initialdir=self.reports_dir,
                initialfile=title.replace(' ', '_') + '.txt')
            if p:
                with open(p, 'w', encoding='utf-8') as f:
                    f.write(txt.get('1.0', tk.END))

        ttk.Button(dlg, text="💾 Сохранить", command=save_as_txt).pack(pady=5)

    def show_charts(self):
        if not HAS_MPL:
            messagebox.showinfo("Диаграммы",
                "Установите matplotlib:\npip install matplotlib")
            return
        dlg = tk.Toplevel(self.root)
        dlg.title("Диаграммы")
        dlg.geometry("900x700")
        dlg.grab_set()
        fig, axes = plt.subplots(2, 2, figsize=(10, 8))

        data = self.filtered()
        acc_sum = {}
        resp_sum = {}
        for a in data:
            acc_sum[a['account']] = acc_sum.get(a['account'], 0) + a['cost']
            resp_sum[a['responsible']] = resp_sum.get(a['responsible'], 0) + a['cost']

        if acc_sum:
            axes[0, 0].pie(list(acc_sum.values()),
                           labels=[x[:20] for x in acc_sum.keys()],
                           autopct='%1.1f%%', textprops={'fontsize': 7})
        axes[0, 0].set_title('Стоимость по счетам')

        top_resp = dict(sorted(resp_sum.items(), key=lambda x: -x[1])[:8])
        if top_resp:
            axes[0, 1].barh(list(top_resp.keys()), list(top_resp.values()))
        axes[0, 1].set_title('Стоимость по ответственным')
        axes[0, 1].tick_params(axis='y', labelsize=7)

        years = {}
        for a in data:
            if a['date']:
                try:
                    y = self._parse_date(a['date']).year
                    years[y] = years.get(y, 0) + a['cost']
                except Exception:
                    pass
        if years:
            axes[1, 0].bar(list(years.keys()), list(years.values()))
            axes[1, 0].tick_params(axis='x', labelsize=8)
        axes[1, 0].set_title('Поступление по годам')

        active = len([a for a in data if not a.get('disposal_date')])
        disposed = len(data) - active
        if active or disposed:
            axes[1, 1].pie([active, disposed], labels=['Активные', 'Выбывшие'],
                           autopct='%1.1f%%', colors=['#4CAF50', '#F44336'])
        axes[1, 1].set_title('Статус')

        plt.tight_layout()
        canvas = FigureCanvasTkAgg(fig, master=dlg)
        canvas.draw()
        canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        plt.close(fig)

    # ================= ПРОВЕРКА / НАПОМИНАНИЯ =================
    def check_integrity(self):
        problems = []
        invs = {}
        for a in self.assets:
            inv = a['inventory'].strip()
            if inv:
                if inv in invs:
                    problems.append(f"Дубликат инв. номера: {inv} "
                                    f"(«{invs[inv]['name']}» и «{a['name']}»)")
                else:
                    invs[inv] = a
            for field in ('account', 'responsible', 'location', 'name'):
                if not a.get(field):
                    problems.append(f"Пустое поле '{field}' у актива "
                                    f"«{a['name'] or a['inventory']}»")
            if a['cost'] < 0:
                problems.append(f"Отрицательная стоимость: {a['name']}")
            if a['depreciation'] < 0:
                problems.append(f"Отрицательная амортизация: {a['name']}")
            if a['depreciation'] > a['cost']:
                problems.append(f"Амортизация > стоимость: {a['name']}")

        if not problems:
            messagebox.showinfo("Проверка", "Ошибок не найдено ✓")
            return

        dlg = tk.Toplevel(self.root)
        dlg.title("Проблемы")
        dlg.geometry("700x400")
        dlg.grab_set()
        txt = tk.Text(dlg, font=('Segoe UI', 10))
        txt.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        for p in problems:
            txt.insert(tk.END, "• " + p + "\n")
        txt.config(state=tk.DISABLED)

    def show_notifications(self):
        messages = []
        today = datetime.now()
        wd = self.show_warranty_days
        for a in self.assets:
            if a.get('disposal_date'):
                continue
            if a.get('warranty_to'):
                try:
                    days = (self._parse_date(a['warranty_to']) - today).days
                    if 0 <= days <= wd:
                        messages.append(
                            f"⚠ Гарантия истекает через {days} дн.: "
                            f"«{a['name']}» (до {a['warranty_to']})")
                except Exception:
                    pass
            if a.get('next_to'):
                try:
                    days = (self._parse_date(a['next_to']) - today).days
                    if 0 <= days <= wd:
                        messages.append(
                            f"🔧 ТО через {days} дн.: «{a['name']}» ({a['next_to']})")
                except Exception:
                    pass
        fully = [a for a in self.assets
                 if a['cost'] > 0 and a['depreciation'] >= a['cost']
                 and not a.get('disposal_date')]
        if fully:
            messages.append(f"💰 Полностью самортизировано: {len(fully)} активов")

        if not messages:
            messagebox.showinfo("Напоминания", "Нет актуальных уведомлений")
            return
        messagebox.showinfo("Напоминания", "\n\n".join(messages[:30]))

    def startup_checks(self):
        self.update_dict_from_assets()
        self.autosave_label.config(
            text=f"Автосохранение каждые {self.autosave_minutes} мин")

    def open_journal(self):
        if not os.path.exists(self.journal_path):
            messagebox.showinfo("Журнал", "Журнал пуст")
            return
        try:
            os.startfile(self.journal_path)
        except Exception:
            pass

    def show_help(self):
        messagebox.showinfo("Горячие клавиши",
            "Ctrl+O  — открыть файл\n"
            "Ctrl+S  — сохранить\n"
            "Ctrl+Shift+S — сохранить как\n"
            "Ctrl+N  — добавить\n"
            "Ctrl+E  — редактировать\n"
            "Delete  — удалить\n"
            "Ctrl+F  — поиск\n"
            "F5      — обновить\n"
            "Двойной клик — карточка\n"
            "Правый клик — контекстное меню")

    # ================= АВТОСОХРАНЕНИЕ =================
    def schedule_autosave(self):
        ms = max(1, self.autosave_minutes) * 60 * 1000
        self.root.after(ms, self.do_autosave)

    def do_autosave(self):
        if self.current_file and not self.current_file.lower().endswith('.xls'):
            self.save_file()
            self.autosave_label.config(
                text=f"✓ Сохранено в {datetime.now().strftime('%H:%M:%S')}")
        self.schedule_autosave()

    def on_close(self):
        self.window_geometry = self.root.geometry()
        self.save_settings()
        if messagebox.askyesno("Выход", "Сохранить изменения перед выходом?"):
            self.save_file()
        self.root.destroy()


# ================= ДИАЛОГ АКТИВА =================
class AssetDialog:
    def __init__(self, parent, title, asset=None, dictionaries=None, last=None):
        self.dialog = tk.Toplevel(parent)
        self.dialog.title(title)
        self.dialog.geometry("540x700")
        self.dialog.configure(bg='#f0f0f0')
        self.dialog.resizable(False, False)
        self.dialog.grab_set()

        self.result = None
        self.asset = asset or {}
        self.dicts = dictionaries or {'accounts': [], 'responsible': [],
                                       'locations': [], 'units': ['шт.']}
        last = last or {}

        tk.Label(self.dialog, text=title, font=('Segoe UI', 14, 'bold'),
                 bg='#f0f0f0').pack(pady=10)

        form = tk.Frame(self.dialog, bg='#f0f0f0')
        form.pack(padx=20, fill=tk.X)

        self.entries = {}
        fields = [
            ('Счет:',            'account',      True,  'accounts'),
            ('Ответственный:',   'responsible',  True,  'responsible'),
            ('Место хранения:',  'location',     True,  'locations'),
            ('Основное средство:', 'name',       True,  None),
            ('Инвентарный номер:', 'inventory',  False, None),
            ('Дата принятия:',    'date',        False, None),
            ('Балансовая стоимость:', 'cost',    False, None),
            ('Количество:',       'quantity',    False, None),
            ('Сумма амортизации:', 'depreciation', False, None),
            ('Гарантия до:',      'warranty_to', False, None),
            ('Следующее ТО:',     'next_to',     False, None),
        ]
        for i, (label, key, req, dict_key) in enumerate(fields):
            tk.Label(form, text=label + (' *' if req else ''),
                     bg='#f0f0f0', font=('Segoe UI', 10)).grid(row=i, column=0,
                                                                sticky='w', pady=3)
            default = self.asset.get(key, '')
            if not default:
                if key == 'quantity':
                    default = 1
                elif key in ('account', 'responsible', 'location'):
                    default = last.get(key, '')
            var = tk.StringVar(value=str(default))
            if dict_key and self.dicts.get(dict_key):
                ttk.Combobox(form, textvariable=var, width=30,
                             values=self.dicts[dict_key]).grid(row=i, column=1,
                                                                sticky='w', padx=8)
            else:
                tk.Entry(form, textvariable=var, width=32,
                         font=('Segoe UI', 10)).grid(row=i, column=1,
                                                      sticky='w', padx=8)
            self.entries[key] = var

        tk.Label(form, text='Остаточная стоимость:', bg='#f0f0f0',
                 font=('Segoe UI', 10)).grid(row=len(fields), column=0,
                                              sticky='w', pady=3)
        self.residual_var = tk.StringVar(value='0.00')
        tk.Label(form, textvariable=self.residual_var, bg='#f0f0f0',
                 font=('Segoe UI', 10, 'bold'),
                 fg='#1a237e').grid(row=len(fields), column=1, sticky='w', padx=8)

        def update_residual(*a):
            try:
                c = float(self.entries['cost'].get().replace(',', '.') or 0)
                d = float(self.entries['depreciation'].get().replace(',', '.') or 0)
                self.residual_var.set(f"{c - d:,.2f} ₽")
            except Exception:
                self.residual_var.set('—')

        for k in ('cost', 'depreciation'):
            self.entries[k].trace('w', update_residual)
        update_residual()

        btns = tk.Frame(self.dialog, bg='#f0f0f0')
        btns.pack(pady=15)
        ttk.Button(btns, text="💾 Сохранить",
                   command=self.save).pack(side=tk.LEFT, padx=8)
        ttk.Button(btns, text="Отмена",
                   command=self.dialog.destroy).pack(side=tk.LEFT, padx=8)

    def save(self):
        for key in ('account', 'responsible', 'location', 'name'):
            if not self.entries[key].get().strip():
                messagebox.showwarning("Внимание", "Заполните обязательные поля (*)")
                return
        try:
            cost = float(self.entries['cost'].get().replace(' ', '').replace(',', '.') or 0)
            dep = float(self.entries['depreciation'].get().replace(' ', '').replace(',', '.') or 0)
            qty = int(self.entries['quantity'].get().strip() or 1)
        except ValueError:
            messagebox.showerror("Ошибка",
                "Стоимость, амортизация и количество должны быть числами")
            return

        self.result = {
            'account':      self.entries['account'].get().strip(),
            'responsible':  self.entries['responsible'].get().strip(),
            'location':     self.entries['location'].get().strip(),
            'name':         self.entries['name'].get().strip(),
            'inventory':    self.entries['inventory'].get().strip(),
            'date':         self.entries['date'].get().strip(),
            'cost':         cost,
            'quantity':     qty,
            'depreciation': dep,
            'disposal_date':   self.asset.get('disposal_date', ''),
            'disposal_reason': self.asset.get('disposal_reason', ''),
            'warranty_to':     self.entries['warranty_to'].get().strip(),
            'next_to':         self.entries['next_to'].get().strip(),
            'residual':        cost - dep,
        }
        self.dialog.destroy()


def main():
    root = tk.Tk()
    AssetManager(root)
    root.mainloop()


if __name__ == "__main__":
    main()
